"""
AR Outstanding Report builder.

Implements the logic in AR_Report_Logic.pdf, validated row-for-row against the
actual "Total Bill Wise Report", "Discharge Wise IP Revenue Report" and a real
"AR as on <date>" output file on 2026-10-12 (Public Ins and TPA & CPR totals,
both old-month and current-month, reproduced to the rupee).

Column selection is by NORMALISED HEADER NAME, not by position. The source
reports are system exports whose header text has inconsistent
whitespace/line-breaks/punctuation between downloads, and whose total column
count and column order can also change between report templates (e.g. a
19-column "Total Bill Wise Report" vs. a newer 13-column one). Each required
field (Date, Category, Company Name, Payment Due, etc.) is located by matching
its header against a small list of accepted spellings (see *_ALIASES below),
after stripping everything but letters/digits and upper-casing. This makes the
loader tolerant of extra/missing/reordered columns, as long as the columns it
actually needs are present under a recognisable header. A column that can't be
found at all raises SourceFormatError naming the missing field, rather than
silently reading the wrong data.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

START_DATE = pd.Timestamp(2021, 4, 1)

# Category routing (validated against real data + user decision on 2026-10-12:
# a Category of ONGC with a real Company Name stays in TPA & CPR, same as any
# other payer -- only a genuinely blank Company Name, or one of the categories
# below, sends a bill to OP GNL).
GNL_PREFIXES = ["GNL", "STUDY-FACULTY", "DEP", "FDM", "P.CARD.FUND", "PM FUND"]
PUBLIC_KEYWORDS = ["ESI", "ECHS", "CHS"]

# Known TPA & CPR category prefixes -- used only to flag unrecognised category
# text for manual review; it never changes where a bill is actually routed.
KNOWN_TPA_PREFIXES = ["TPA", "CPR", "LABCPR", "CSLCPR", "ONGC", "INHS", "SSV-TPA", "EMP-TPA"]

OUT_COLUMNS = [
    "SlNo", "BillNo", "Date", "MRD", "VisitCode", "BillType", "PatientName",
    "Category", "TotalAmt", "NetAmt", "PaidAmt", "PaymentDue", "CompanyName",
]

DETAIL_HEADERS_WITH_CATEGORY = [
    "Sl.No.", "Bill No.", "Date", "MRD No.", "Visit Code", "Bill Type",
    "Patient Name", "Category", "Total Amt (Inc. Tax)", "Net Amt", "Paid Amt",
    "Payment Due", "Company Name",
]
DETAIL_HEADERS_NO_CATEGORY = [h for h in DETAIL_HEADERS_WITH_CATEGORY if h != "Category"]

MONEY_COLS_WITH_CATEGORY = [9, 10, 11, 12]   # 1-indexed columns of TotalAmt..PaymentDue
MONEY_COLS_NO_CATEGORY = [8, 9, 10, 11]

# Accepted header spellings per logical field, after _normalize_header() strips
# everything but letters/digits and upper-cases. Listed in preference order;
# the first one found in the file wins. Add a new spelling here -- never a new
# fixed column position -- when a report template renames a column.
OP_ALIASES: dict[str, list[str]] = {
    "SlNo": ["SLNO"],
    "BillNo": ["BILLNO"],
    "Date": ["DATE", "BILLDATE", "DISCHARGEDATE"],
    "MRD": ["MRDNO", "MRD"],
    "VisitCode": ["VISITCODE"],
    "BillType": ["BILLTYPE"],
    "PatientName": ["PATIENTNAME"],
    "Category": ["CATEGORY"],
    "TotalAmt": ["TOTALAMTINCTAX", "TOTALAMT"],
    "NetAmt": ["NETAMT"],
    "PaidAmt": ["PAIDAMT"],
    "PaymentDue": ["PAYMENTDUE"],
    "CompanyName": ["COMPANYNAME"],
}

IP_ALIASES: dict[str, list[str]] = {
    "SlNo": ["SLNO"],
    "BillNo": ["BILLNO"],
    "Date": ["DISCHARGEDATE", "DATE", "BILLDATE"],
    "MRD": ["MRDNO", "MRD"],
    "VisitCode": ["VISITCODE"],
    "BillType": ["BILLTYPE"],
    "PatientName": ["PATIENTNAME"],
    "Category": ["CATEGORY"],
    "TotalAmt": ["TOTALAMTINCTAX", "TOTALAMT"],
    "NetAmt": ["NETAMT"],
    "PaidAmt": ["PAIDAMT"],
    "PaymentDue": ["PAYMENTDUE"],
    "CompanyName": ["COMPANYNAME"],
    "BillStatus": ["BILLSTATUS"],
    "ConfirmedDate": ["CONFIRMEDDATE"],
}

# IP-only columns used purely for the shifted-row fix, in on-disk left-to-right
# order starting at Bill Status. Only BillStatus/ConfirmedDate need aliases
# (above) to detect the shift; the rest just need to exist so their values can
# be carried across when fixing it, so any header is accepted positionally
# relative to Bill Status/Confirmed Date.
IP_SHIFT_FIELD_COUNT = 8  # Bill Status, Confirmed Date, Admission Date, No. of days Admitted, Department, Doctor, Billing Type, Company Name

# Human-readable labels for the field keys above, for a UI to show someone
# mapping a field to one of the file's actual columns.
FIELD_LABELS: dict[str, str] = {
    "SlNo": "Sl. No.",
    "BillNo": "Bill No.",
    "Date": "Date (Bill Date / Discharge Date)",
    "MRD": "MRD No.",
    "VisitCode": "Visit Code",
    "BillType": "Bill Type",
    "PatientName": "Patient Name",
    "Category": "Category",
    "TotalAmt": "Total Amt (Inc. Tax)",
    "NetAmt": "Net Amt",
    "PaidAmt": "Paid Amt",
    "PaymentDue": "Payment Due",
    "CompanyName": "Company Name",
    "BillStatus": "Bill Status",
    "ConfirmedDate": "Confirmed Date",
}


class SourceFormatError(ValueError):
    """Raised when a source file doesn't match what this tool expects."""


class MissingColumnsError(SourceFormatError):
    """Raised when one or more required columns can't be located in a source
    file by header name, even after any manual column-name mapping supplied.
    Carries structured data, not just a message, so a UI can offer to map
    each missing field to one of the file's actual headers."""

    def __init__(self, report_label: str, missing_fields: list[str], available_headers: list[str]):
        self.report_label = report_label
        self.missing_fields = missing_fields
        self.available_headers = available_headers
        wanted = ", ".join(FIELD_LABELS.get(f, f) for f in missing_fields)
        super().__init__(
            f"{report_label}: couldn't find a column for: {wanted}.\n"
            f"The file's actual headers are: {available_headers}\n"
            "The report's layout may have changed -- map each missing field to "
            "the right column before re-running, rather than trusting the output."
        )


@dataclass
class BuildResult:
    output_path: Path
    report_date: pd.Timestamp
    bill_count_public: int
    bill_count_tpa: int
    bill_count_gnl: int
    public_total: float
    tpa_total: float
    needs_review: pd.DataFrame
    previous_report_date: str | None
    warnings: list[str] = field(default_factory=list)


def _normalize_header(header) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(header).upper())


def _locate_fields(df: pd.DataFrame, aliases: dict[str, list[str]], report_label: str,
                    manual_map: dict[str, str] | None = None) -> dict[str, int]:
    """Locate every field in `aliases` as a 0-based column position in `df`.
    `manual_map` (field -> the file's actual header text, as a human chose it
    in the app's column-mapping dialog) is tried first for each field, ahead
    of the automatic alias list. Every missing field is collected and raised
    together as one MissingColumnsError, instead of stopping at the first."""
    normalized = [_normalize_header(c) for c in df.columns]
    cols: dict[str, int] = {}
    missing: list[str] = []
    for field, opts in aliases.items():
        idx = None
        manual_header = (manual_map or {}).get(field)
        if manual_header is not None:
            manual_norm = _normalize_header(manual_header)
            idx = next((i for i, n in enumerate(normalized) if n == manual_norm), None)
        if idx is None:
            for alias in opts:
                idx = next((i for i, n in enumerate(normalized) if n == alias), None)
                if idx is not None:
                    break
        if idx is None:
            missing.append(field)
        else:
            cols[field] = idx
    if missing:
        raise MissingColumnsError(report_label, missing, list(df.columns))
    return cols


def _read_source_table(path: str | Path) -> pd.DataFrame:
    """Read a source report, whichever file format it was exported as."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, header=0, low_memory=False)
    return pd.read_excel(path, sheet_name=0, header=0)


def _as_money(series: pd.Series) -> pd.Series:
    """Coerce an amount column to numeric. Source exports occasionally embed
    stray summary/label rows whose Bill No. is blank (and get dropped), but
    while present they can force an otherwise-numeric column to text dtype
    for the whole column -- left uncorrected, summing it later would silently
    do string concatenation instead of addition."""
    return pd.to_numeric(series, errors="coerce")


def load_op_bills(path: str | Path, column_overrides: dict[str, str] | None = None) -> pd.DataFrame:
    df = _read_source_table(path)
    cols = _locate_fields(df, OP_ALIASES, "Total Bill Wise Report", column_overrides)
    df = df[df.iloc[:, cols["BillNo"]].notna()].copy()  # drop any row with no Bill No.

    out = pd.DataFrame({
        "SlNo": df.iloc[:, cols["SlNo"]],
        "BillNo": df.iloc[:, cols["BillNo"]],
        "Date": pd.to_datetime(df.iloc[:, cols["Date"]]),
        "MRD": df.iloc[:, cols["MRD"]],
        "VisitCode": df.iloc[:, cols["VisitCode"]],
        "BillType": df.iloc[:, cols["BillType"]],
        "PatientName": df.iloc[:, cols["PatientName"]],
        "Category": df.iloc[:, cols["Category"]],
        "TotalAmt": _as_money(df.iloc[:, cols["TotalAmt"]]),
        "NetAmt": _as_money(df.iloc[:, cols["NetAmt"]]),
        "PaidAmt": _as_money(df.iloc[:, cols["PaidAmt"]]),
        "PaymentDue": _as_money(df.iloc[:, cols["PaymentDue"]]),
        "CompanyName": df.iloc[:, cols["CompanyName"]],
    })
    out["Source"] = "OP"
    return out.reset_index(drop=True)


def load_ip_bills(path: str | Path, column_overrides: dict[str, str] | None = None) -> pd.DataFrame:
    df = _read_source_table(path)
    cols = _locate_fields(df, IP_ALIASES, "Discharge Wise IP Revenue Report", column_overrides)
    df = df[df.iloc[:, cols["BillNo"]].notna()].copy()  # drop heading-counter rows and grand-total footer rows
    df = df.reset_index(drop=True)

    # Fix rows where every column from Bill Status onward got shifted one
    # place right, so "CONF"/"OPN" shows up under Confirmed Date instead.
    bs_idx, cd_idx = cols["BillStatus"], cols["ConfirmedDate"]
    if cd_idx != bs_idx + 1 or cols["CompanyName"] != bs_idx + IP_SHIFT_FIELD_COUNT - 1:
        raise SourceFormatError(
            "Discharge Wise IP Revenue Report: the columns from 'Bill Status' to "
            "'Company Name' are no longer in the expected consecutive order, so "
            "the shifted-row fix can't run safely. Check the file's layout "
            "before re-running."
        )
    shift_idxs = list(range(bs_idx, bs_idx + IP_SHIFT_FIELD_COUNT))  # Bill Status .. Company Name, on-disk order
    shifted = df.iloc[:, cd_idx].isin(["CONF", "OPN"])
    n_shifted = int(shifted.sum())
    if n_shifted:
        ncols = df.shape[1]
        src_cols = [
            df.iloc[:, i].copy() if i < ncols else pd.Series([pd.NA] * len(df), index=df.index)
            for i in [j + 1 for j in shift_idxs]
        ]  # each field's value, one column to the right of where it belongs -- snapshot before overwriting
        for offset, i in enumerate(shift_idxs):
            df.loc[shifted, df.columns[i]] = src_cols[offset][shifted].values

    df = df[(df.iloc[:, cols["BillType"]] == "IP_F") & (df.iloc[:, bs_idx] == "CONF")].copy()

    out = pd.DataFrame({
        "SlNo": df.iloc[:, cols["SlNo"]],
        "BillNo": df.iloc[:, cols["BillNo"]],
        "Date": pd.to_datetime(df.iloc[:, cols["Date"]]),  # Discharge Date -> Date
        "MRD": df.iloc[:, cols["MRD"]],
        "VisitCode": df.iloc[:, cols["VisitCode"]],
        "BillType": df.iloc[:, cols["BillType"]],
        "PatientName": df.iloc[:, cols["PatientName"]],
        "Category": df.iloc[:, cols["Category"]],
        "TotalAmt": _as_money(df.iloc[:, cols["TotalAmt"]]),
        "NetAmt": _as_money(df.iloc[:, cols["NetAmt"]]),
        "PaidAmt": _as_money(df.iloc[:, cols["PaidAmt"]]),
        "PaymentDue": _as_money(df.iloc[:, cols["PaymentDue"]]),
        "CompanyName": df.iloc[:, cols["CompanyName"]],
    })
    out["Source"] = "IP"
    out.attrs["shifted_rows_fixed"] = n_shifted
    return out.reset_index(drop=True)


def _normalize_category(category) -> str:
    return str(category).strip().upper()


def _normalize_company(company_name) -> str:
    return "" if pd.isna(company_name) else str(company_name).strip().upper()


def override_key(category, company_name) -> str:
    """Two bills can share the exact same Category text but belong on
    different sheets depending on who's actually paying (e.g. Category
    'STUDY' with Company Name 'ESI- Beneficiary ...' vs. 'Aasandha Company
    Limited') -- so a human reviewer's decision, and the override key it's
    stored under, is per (Category, Company Name) pair, not per Category
    alone."""
    return f"{_normalize_category(category)}|{_normalize_company(company_name)}"


def route_sheet(category, company_name, overrides: dict[str, str] | None = None) -> str:
    """Decide which sheet a bill belongs on. `overrides` maps an
    override_key(category, company_name) to a bucket ("OP GNL" / "PUBLIC" /
    "TPA"), as decided by a human reviewing the NEEDS REVIEW list in the app
    -- it takes priority over every rule below, including the blank-company
    GNL default."""
    if overrides:
        key = override_key(category, company_name)
        if key in overrides:
            return overrides[key]
    cat = _normalize_category(category)
    if pd.isna(company_name):
        return "OP GNL"
    if any(cat.startswith(p) for p in GNL_PREFIXES):
        return "OP GNL"
    if any(k in cat for k in PUBLIC_KEYWORDS):
        return "PUBLIC"
    return "TPA"


def _is_known_tpa_category(category, company_name=None, overrides: dict[str, str] | None = None) -> bool:
    if overrides and override_key(category, company_name) in overrides:
        return True  # a human already made an explicit call on this category/company pair
    cat = _normalize_category(category)
    return any(cat.startswith(p) for p in KNOWN_TPA_PREFIXES)


def build_combined(op_df: pd.DataFrame, ip_df: pd.DataFrame, report_date: pd.Timestamp,
                    category_overrides: dict[str, str] | None = None) -> pd.DataFrame:
    combined = pd.concat([op_df, ip_df], ignore_index=True)
    end = report_date + pd.Timedelta(hours=23, minutes=59, seconds=59)
    combined = combined[(combined["Date"] >= START_DATE) & (combined["Date"] <= end)].copy()
    combined["Sheet"] = [
        route_sheet(c, n, category_overrides) for c, n in zip(combined["Category"], combined["CompanyName"])
    ]
    return combined.reset_index(drop=True)


def needs_review_rows(combined: pd.DataFrame, category_overrides: dict[str, str] | None = None) -> pd.DataFrame:
    tpa = combined[combined["Sheet"] == "TPA"]
    known = tpa.apply(
        lambda r: _is_known_tpa_category(r["Category"], r["CompanyName"], category_overrides), axis=1
    )
    unknown = tpa[~known]
    return unknown[["BillNo", "Date", "Category", "CompanyName", "PaymentDue", "Source"]].copy()


def compute_table1(combined: pd.DataFrame, report_date: pd.Timestamp) -> dict:
    month_start = pd.Timestamp(report_date.year, report_date.month, 1)
    end = report_date + pd.Timedelta(hours=23, minutes=59, seconds=59)
    old = combined[combined["Date"] < month_start]
    cur = combined[(combined["Date"] >= month_start) & (combined["Date"] <= end)]

    result = {"month_start": month_start}
    for key, label in (("PUBLIC", "public"), ("TPA", "tpa")):
        o = float(old[old["Sheet"] == key]["PaymentDue"].sum())
        c = float(cur[cur["Sheet"] == key]["PaymentDue"].sum())
        result[f"{label}_old"] = o
        result[f"{label}_cur"] = c
        result[f"{label}_total"] = o + c
    result["total_old"] = result["public_old"] + result["tpa_old"]
    result["total_cur"] = result["public_cur"] + result["tpa_cur"]
    result["total_total"] = result["public_total"] + result["tpa_total"]
    return result


def _parse_title_date(title) -> pd.Timestamp | None:
    m = re.search(r"as on\s+(\d{1,2}[/-]\d{1,2}[/-]\d{4})", str(title or ""))
    if not m:
        return None
    try:
        return pd.to_datetime(m.group(1), dayfirst=True)
    except (ValueError, TypeError):
        return None


def read_previous_totals(previous_path: str | Path) -> dict | None:
    """Read Table 1's second column (the pre-current-month "old" bucket) for
    Public/TPA from a prior AR output file, and the report date it was
    generated for (parsed from its title row). This is the column to compare
    against this run's own old bucket in Table 2, since both describe the
    same date range when the previous report falls in the same month as this
    one."""
    wb = load_workbook(previous_path, data_only=True)
    if "OUTSTANDING STATUS" not in wb.sheetnames:
        return None
    ws = wb["OUTSTANDING STATUS"]

    prev_date = _parse_title_date(ws.cell(row=1, column=1).value)

    public_old = tpa_old = None
    for row in ws.iter_rows(min_row=1, max_row=15, values_only=False):
        label = row[0].value
        if label == "OP&IP Public confirmed" and public_old is None:
            public_old = row[1].value
        elif label == "OP&IP TPI confirmed" and tpa_old is None:
            tpa_old = row[1].value
        if public_old is not None and tpa_old is not None:
            break

    if public_old is None or tpa_old is None:
        return None
    return {
        "previous_report_date": prev_date,
        "public_old": float(public_old),
        "tpa_old": float(tpa_old),
    }


def compute_table2(table1: dict, previous: dict | None) -> dict | None:
    if previous is None:
        return None
    return {
        "previous_report_date": previous["previous_report_date"],
        "public_prev": previous["public_old"],
        "tpa_prev": previous["tpa_old"],
        "public_cur": table1["public_old"],
        "tpa_cur": table1["tpa_old"],
        "public_decrease": previous["public_old"] - table1["public_old"],
        "tpa_decrease": previous["tpa_old"] - table1["tpa_old"],
    }


def _money_format(ws, row: int, cols: list[int]) -> None:
    for c in cols:
        ws.cell(row=row, column=c).number_format = "#,##0.00"


def _write_detail_sheet(wb: Workbook, name: str, df: pd.DataFrame, report_date: pd.Timestamp,
                         with_title: bool, with_category: bool) -> None:
    ws = wb.create_sheet(name)
    headers = DETAIL_HEADERS_WITH_CATEGORY if with_category else DETAIL_HEADERS_NO_CATEGORY
    money_cols = MONEY_COLS_WITH_CATEGORY if with_category else MONEY_COLS_NO_CATEGORY

    row_idx = 1
    if with_title:
        ws.cell(row=1, column=1,
                 value=f"Accounts receivables during the period {START_DATE:%d/%m/%Y} - {report_date:%d/%m/%Y}")
        ws.cell(row=1, column=1).font = Font(bold=True)
        row_idx = 2

    header_row = row_idx
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=c, value=h)
        cell.font = Font(bold=True)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate

    cols_to_write = ["SlNo", "BillNo", "Date", "MRD", "VisitCode", "BillType", "PatientName"]
    if with_category:
        cols_to_write.append("Category")
    cols_to_write += ["TotalAmt", "NetAmt", "PaidAmt", "PaymentDue", "CompanyName"]

    date_col = cols_to_write.index("Date") + 1
    r = header_row + 1
    for _, rec in df.iterrows():
        for c, col in enumerate(cols_to_write, start=1):
            ws.cell(row=r, column=c, value=rec[col])
        ws.cell(row=r, column=date_col).number_format = "dd-mm-yyyy"
        _money_format(ws, r, money_cols)
        r += 1

    total_row = r + 1
    total_col = money_cols[-1]  # Payment Due column
    ws.cell(row=total_row, column=total_col, value=float(df["PaymentDue"].sum()))
    ws.cell(row=total_row, column=total_col).number_format = "#,##0.00"
    ws.cell(row=total_row, column=total_col).font = Font(bold=True)

    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16


def _write_status_sheet(wb: Workbook, table1: dict, table2: dict | None, report_date: pd.Timestamp) -> None:
    ws = wb.create_sheet("OUTSTANDING STATUS", 0)
    bold = Font(bold=True)

    prev_month_end = table1["month_start"] - pd.Timedelta(days=1)

    ws.cell(row=1, column=1, value=f"Outstanding status as on {report_date:%d/%m/%Y}").font = bold

    ws.cell(row=3, column=1, value="AR")
    ws.cell(row=3, column=2, value=f"{START_DATE:%d/%m/%Y} to {prev_month_end:%d/%m/%Y}")
    ws.cell(row=3, column=3, value=f"{table1['month_start']:%d/%m/%Y} to {report_date:%d/%m/%Y}")
    ws.cell(row=3, column=4, value="Total")
    for c in range(1, 5):
        ws.cell(row=3, column=c).font = bold
    ws.cell(row=4, column=2, value="Rs")

    rows = [
        ("OP&IP Public confirmed", table1["public_old"], table1["public_cur"], table1["public_total"]),
        ("OP&IP TPI confirmed", table1["tpa_old"], table1["tpa_cur"], table1["tpa_total"]),
        ("Total receivables", table1["total_old"], table1["total_cur"], table1["total_total"]),
    ]
    r = 5
    for label, old, cur, total in rows:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=old).number_format = "#,##0.00"
        ws.cell(row=r, column=3, value=cur).number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=total).number_format = "#,##0.00"
        if label == "Total receivables":
            for c in range(1, 5):
                ws.cell(row=r, column=c).font = bold
        r += 1

    if table2 is not None:
        r += 2
        ws.cell(row=r, column=1, value=f"Decrease in Outstanding status as on {report_date:%d/%m/%Y}").font = bold
        r += 2
        ws.cell(row=r, column=1, value="AR")
        ws.cell(row=r, column=2, value=f"{START_DATE:%d-%m-%Y} to {prev_month_end:%d-%m-%Y}")
        r += 1
        prd = table2["previous_report_date"]
        if prd is None:
            prev_label = "as on previous report"
        elif prd == report_date:
            prev_label = "as on previous report (same date)"
        else:
            prev_label = f"as on {prd:%d/%m/%Y}"
        ws.cell(row=r, column=2, value=prev_label)
        ws.cell(row=r, column=3, value=f"as on {report_date:%d/%m/%Y}")
        ws.cell(row=r, column=4, value="Decrease")
        for c in range(1, 5):
            ws.cell(row=r, column=c).font = bold
        r += 1
        ws.cell(row=r, column=2, value="Rs")
        r += 1

        rows2 = [
            ("OP&IP Public confirmed", table2["public_prev"], table2["public_cur"], table2["public_decrease"]),
            ("OP&IP TPI confirmed", table2["tpa_prev"], table2["tpa_cur"], table2["tpa_decrease"]),
        ]
        total_prev = table2["public_prev"] + table2["tpa_prev"]
        total_cur = table2["public_cur"] + table2["tpa_cur"]
        total_dec = table2["public_decrease"] + table2["tpa_decrease"]
        rows2.append(("Total receivables", total_prev, total_cur, total_dec))
        for label, prev, cur, dec in rows2:
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=prev).number_format = "#,##0.00"
            ws.cell(row=r, column=3, value=cur).number_format = "#,##0.00"
            ws.cell(row=r, column=4, value=dec).number_format = "#,##0.00"
            if label == "Total receivables":
                for c in range(1, 5):
                    ws.cell(row=r, column=c).font = bold
            r += 1

    for c, w in zip(range(1, 5), [24, 22, 22, 18]):
        ws.column_dimensions[get_column_letter(c)].width = w


def _write_needs_review_sheet(wb: Workbook, df: pd.DataFrame) -> None:
    if df.empty:
        return
    ws = wb.create_sheet("NEEDS REVIEW")
    ws.cell(row=1, column=1,
            value="Bills routed to TPA & CPR with a category this tool doesn't recognise -- please check.").font = Font(bold=True)
    headers = ["Bill No.", "Date", "Category", "Company Name", "Payment Due", "Source"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=2, column=c, value=h).font = Font(bold=True)
    r = 3
    for _, rec in df.iterrows():
        ws.cell(row=r, column=1, value=rec["BillNo"])
        ws.cell(row=r, column=2, value=rec["Date"]).number_format = "dd-mm-yyyy"
        ws.cell(row=r, column=3, value=rec["Category"])
        ws.cell(row=r, column=4, value=rec["CompanyName"])
        ws.cell(row=r, column=5, value=rec["PaymentDue"]).number_format = "#,##0.00"
        ws.cell(row=r, column=6, value=rec["Source"])
        r += 1
    for c, w in zip(range(1, 7), [24, 14, 20, 30, 14, 10]):
        ws.column_dimensions[get_column_letter(c)].width = w


def _check_op_bill_types(op_df: pd.DataFrame) -> None:
    """Guard against the wrong file being supplied as the OP report -- e.g. an
    IP/Discharge export relabelled or re-used as "Total Bill Wise Report".
    OP bills use types like O_B; IP bills use types like IP_F."""
    if op_df.empty:
        return
    bill_type = op_df["BillType"].astype(str).str.upper()
    ip_like_fraction = bill_type.str.contains("IP").mean()
    if ip_like_fraction > 0.5:
        raise SourceFormatError(
            "The OP file's 'Bill Type' column is mostly IP-style values (e.g. "
            "IP_F), not OP values like O_B. This looks like an IP/Discharge "
            "report supplied as the OP 'Total Bill Wise Report' -- please "
            "check which file you selected before re-running."
        )


def infer_report_date(op_df: pd.DataFrame, ip_df: pd.DataFrame) -> pd.Timestamp:
    """The report date is whatever date the two downloads run through --
    i.e. the latest Bill Date / Discharge Date present in either file."""
    return max(op_df["Date"].max(), ip_df["Date"].max()).normalize()


def build_report(op_path: str | Path, ip_path: str | Path,
                  previous_path: str | Path | None = None,
                  output_dir: str | Path | None = None,
                  category_overrides: dict[str, str] | None = None,
                  op_column_overrides: dict[str, str] | None = None,
                  ip_column_overrides: dict[str, str] | None = None) -> BuildResult:
    """Build the AR Outstanding Report from the two source files, plus an
    optional prior AR report. The prior report, when given, supplies Table
    2's Decrease column (its own Table 1 old-bucket figures for Public/TPA).
    Without it, the OUTSTANDING STATUS sheet holds just Table 1 -- the
    current outstanding totals -- and Table 2 is skipped. The report date is
    inferred from the data, and the output folder defaults to the OP file's
    folder.

    `category_overrides` maps an override_key(category, company_name) to the
    bucket a human reviewer has decided that pair belongs on ("OP GNL" /
    "PUBLIC" / "TPA") -- it overrides the default routing rules for that
    exact category+company pair and stops it being flagged on the NEEDS
    REVIEW sheet again. Keying on the pair, not the category alone, matters
    because two bills can share the same Category text (e.g. "STUDY") while
    actually belonging to different payers that need different decisions.

    `op_column_overrides` / `ip_column_overrides` map a field name (e.g.
    "PaymentDue", see FIELD_LABELS) to the exact header text a human has
    confirmed that field actually lives under in that file -- used when a
    report template renames a column to something _locate_fields' built-in
    alias list doesn't recognise, raising MissingColumnsError. Like
    `category_overrides`, the app persists these across runs; this function
    itself has no opinion on where they're stored."""
    warnings: list[str] = []

    op_path = Path(op_path)
    previous_path = Path(previous_path) if previous_path else None
    output_dir = Path(output_dir) if output_dir else op_path.parent
    output_dir.mkdir(parents=True, exist_ok=True)

    op_df = load_op_bills(op_path, op_column_overrides)
    _check_op_bill_types(op_df)
    ip_df = load_ip_bills(ip_path, ip_column_overrides)
    n_shifted = ip_df.attrs.get("shifted_rows_fixed", 0)
    if n_shifted:
        warnings.append(f"Fixed {n_shifted} shifted row(s) in the IP file (Bill Status column misaligned).")

    report_date = infer_report_date(op_df, ip_df)

    combined = build_combined(op_df, ip_df, report_date, category_overrides)
    review_df = needs_review_rows(combined, category_overrides)
    if not review_df.empty:
        warnings.append(f"{len(review_df)} bill(s) had an unrecognised TPA & CPR category -- see NEEDS REVIEW sheet.")

    table1 = compute_table1(combined, report_date)
    out_path = output_dir / f"AR as on {report_date:%d-%m-%Y}.xlsx"

    if previous_path is None:
        table2 = None
        previous_report_date_str = None
        warnings.append("No previous report supplied -- showing only the current outstanding totals (Decrease table skipped).")
    else:
        previous_totals = read_previous_totals(previous_path)
        if previous_totals is None:
            raise SourceFormatError(
                f"Could not read Table 1 figures from the previous report file "
                f"'{previous_path.name}' -- it must contain an OUTSTANDING STATUS "
                f"sheet with the standard Table 1 layout."
            )
        if previous_totals["previous_report_date"] is None:
            warnings.append("Could not parse the previous report's date from its title row -- shown as 'previous report'.")
        warnings.append(f"Using {previous_path.name} for the Decrease table.")
        prd = previous_totals["previous_report_date"]
        previous_report_date_str = f"{prd:%d/%m/%Y}" if prd is not None else None
        table2 = compute_table2(table1, previous_totals)

    wb = Workbook()
    wb.remove(wb.active)
    _write_status_sheet(wb, table1, table2, report_date)
    _write_detail_sheet(wb, "OP&IP PUBLIC INS", combined[combined["Sheet"] == "PUBLIC"], report_date,
                        with_title=True, with_category=True)
    _write_detail_sheet(wb, "OP&IP TPA & CPR", combined[combined["Sheet"] == "TPA"], report_date,
                        with_title=True, with_category=True)
    _write_detail_sheet(wb, "OP GNL", combined[combined["Sheet"] == "OP GNL"], report_date,
                        with_title=False, with_category=False)
    _write_needs_review_sheet(wb, review_df)
    wb.save(out_path)

    return BuildResult(
        output_path=out_path,
        report_date=report_date,
        bill_count_public=int((combined["Sheet"] == "PUBLIC").sum()),
        bill_count_tpa=int((combined["Sheet"] == "TPA").sum()),
        bill_count_gnl=int((combined["Sheet"] == "OP GNL").sum()),
        public_total=table1["public_total"],
        tpa_total=table1["tpa_total"],
        needs_review=review_df,
        previous_report_date=previous_report_date_str,
        warnings=warnings,
    )
